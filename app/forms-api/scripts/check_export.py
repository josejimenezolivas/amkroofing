"""Exercise the Export menu end to end and report what came out.

Drives the real app in a browser -- opening a document, clicking Export and
picking a format -- so this checks the button wiring, not just the endpoint.
The scratch document is removed afterwards.

    .venv/bin/python scripts/check_export.py invoice
"""

import asyncio
import sys
import zipfile
from pathlib import Path

from playwright.async_api import async_playwright

from devapi import APP, scratch

OUT = Path(__file__).resolve().parents[1] / ".compare"

NEW_DOC = {"invoice": "Invoice", "agreement": "Agreement"}


async def run(template: str) -> None:
    OUT.mkdir(exist_ok=True)

    async with async_playwright() as p:
        browser = await p.chromium.launch()
        page = await browser.new_page(accept_downloads=True)
        await page.goto(APP, wait_until="networkidle")
        await page.get_by_role("button", name=NEW_DOC[template], exact=True).click()
        await page.wait_for_selector(".document .sheet, .document .msheet")

        for label, suffix in [("PDF document", "pdf"), ("Microsoft Word", "docx")]:
            await page.get_by_role("button", name="Export").click()
            async with page.expect_download() as pending:
                await page.get_by_role("menuitem", name=label).click()
            download = await pending.value

            target = OUT / f"export-{template}.{suffix}"
            await download.save_as(target)
            describe(target, download.suggested_filename)

        await browser.close()


def describe(path: Path, suggested: str) -> None:
    size = path.stat().st_size

    if path.suffix == ".pdf":
        import pymupdf

        note = f"{pymupdf.open(path).page_count} pages"
    else:
        from docx import Document

        document = Document(path)
        with zipfile.ZipFile(path) as archive:
            images = [n for n in archive.namelist() if n.startswith("word/media/")]
        note = (f"{len(document.tables)} tables, "
                f"{len(document.paragraphs)} paragraphs, {len(images)} images")

    print(f"  {suggested:<40} {size / 1024:7.1f} KB  {note}")


if __name__ == "__main__":
    for name in sys.argv[1:] or ["invoice", "agreement"]:
        print(name)
        with scratch():
            asyncio.run(run(name))
